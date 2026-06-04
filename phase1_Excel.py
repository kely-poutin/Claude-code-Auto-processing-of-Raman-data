#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Phase 1 - Import TXT spectral data into Excel and organize.
==========================================================
Auto-detects the script's own directory, filters .txt files whose names
contain digits, sorts by the first number found, imports each .txt into
its own worksheet, creates a "Data Collection" sheet, and cleans up.

Suitable for: BWSpec spectrometer .txt exports (UTF-8, semicolon-delimited).
              Raman Shift = column 4, Raw data #1 = column 7.
              Header row = line 79, data starts at line 80.

Usage: Place this .py file in the target folder, then run:
       python phase1.py
"""

import os
import re
import sys
import openpyxl
from openpyxl import Workbook

# ---- Fix Windows GBK stdout for scripts with CJK paths ----
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# ---- Auto-detect script directory ----
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
os.chdir(BASE_DIR)

FOLDER_NAME = os.path.basename(BASE_DIR)
XLSX_NAME = FOLDER_NAME + ".xlsx"
XLSX_PATH = os.path.join(BASE_DIR, XLSX_NAME)

# ---- BWSpec file format constants ----
HEADER_ROW = 79        # 1-based line number of column headers
DATA_START_ROW = 80    # 1-based line number where data begins
RAMAN_SHIFT_COL = 4    # 1-based column index of Raman Shift
RAW_DATA_COL = 7       # 1-based column index of Raw data #1
ENCODING = "utf-8"

print(f"Working dir: {BASE_DIR}")
print(f"Target XLSX: {XLSX_NAME}")
print()

# ---- Step 1: Find .txt files with digits in name, sort by first number ----
all_files = os.listdir(BASE_DIR)
txt_files = []
for f in all_files:
    name_no_ext = os.path.splitext(f)[0]
    if f.lower().endswith(".txt") and re.search(r"\d", name_no_ext):
        txt_files.append(f)

if not txt_files:
    print("[ERR] No .txt files with digits in filename found.")
    print("      Make sure BWSpec .txt files (e.g. 1.txt, sample2.txt) are in this folder.")
    sys.exit(1)


def extract_first_number(filename: str) -> int:
    """Extract the first contiguous digit group from a filename for sorting."""
    m = re.search(r"(\d+)", filename)
    return int(m.group(1)) if m else 999999


txt_files.sort(key=extract_first_number)
print(f"Found {len(txt_files)} matching .txt files:")
for fn in txt_files:
    print(f"  -> {fn}")
print()

# ---- Create Excel workbook, import each .txt ----
wb = Workbook()
default_sheet = wb.active
saved_sheets = []

# Characters forbidden in Excel sheet names
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")

for idx, txt_file in enumerate(txt_files):
    raw_name = os.path.splitext(txt_file)[0]
    safe_name = INVALID_SHEET_CHARS.sub("_", raw_name)[:31]  # Excel 31-char limit

    # Use default sheet for the first file, create new sheets for others
    ws = default_sheet if idx == 0 else wb.create_sheet()
    ws.title = safe_name

    txt_path = os.path.join(BASE_DIR, txt_file)
    with open(txt_path, "r", encoding=ENCODING, errors="replace") as fh:
        lines = fh.readlines()

    if len(lines) < HEADER_ROW:
        print(f"  [WARN] Skipping {txt_file}: fewer than {HEADER_ROW} lines.")
        continue

    # Write column headers (row 1)
    headers = lines[HEADER_ROW - 1].strip().split(";")
    for ci, h in enumerate(headers, 1):
        ws.cell(row=1, column=ci, value=h.strip())

    # Write data (starting from row 2)
    data_lines = lines[DATA_START_ROW - 1:]
    row_count = 0
    for ri, line in enumerate(data_lines):
        line = line.strip()
        if not line:
            continue
        values = line.split(";")
        for ci, val in enumerate(values, 1):
            val = val.strip()
            if val == "":
                ws.cell(row=ri + 2, column=ci, value=None)
            else:
                try:
                    ws.cell(row=ri + 2, column=ci, value=float(val))
                except ValueError:
                    ws.cell(row=ri + 2, column=ci, value=val)
        row_count += 1

    saved_sheets.append(safe_name)
    print(f"  [OK] {txt_file} -> sheet '{safe_name}' ({row_count} data rows)")

if not saved_sheets:
    print("[ERR] No worksheets were created. Aborting.")
    sys.exit(1)

print(f"\nSheets created: {saved_sheets}")

# ---- Step 2: Create "Data Collection" sheet ----
ws_data = wb.create_sheet(title="数据集合")  # 数据集合

first_sheet = wb[saved_sheets[0]]
max_row = first_sheet.max_row
print(f"\nFirst sheet max row: {max_row}")

# Header row
ws_data.cell(row=1, column=1, value="Raman Shift")
for ci, sn in enumerate(saved_sheets):
    ws_data.cell(row=1, column=ci + 2, value=f"{sn} Raw data #1")

# Raman Shift -> Column A (starting row 2)
for ri in range(2, max_row + 1):
    val = first_sheet.cell(row=ri, column=RAMAN_SHIFT_COL).value
    ws_data.cell(row=ri, column=1, value=val)

# Raw data #1 -> Columns B, C, D, ... (starting row 2)
for ci, sn in enumerate(saved_sheets):
    src_ws = wb[sn]
    src_max = src_ws.max_row
    for ri in range(2, src_max + 1):
        val = src_ws.cell(row=ri, column=RAW_DATA_COL).value
        ws_data.cell(row=ri, column=ci + 2, value=val)
    print(f"  [OK] '{sn}' Raw data #1 -> column {ci + 2}")

# ---- Step 3: Delete rows where Column A is empty (bottom-up) ----
max_data_row = ws_data.max_row
print(f"\nData collection rows (before cleanup): {max_data_row}")

rows_to_delete = []
for ri in range(2, max_data_row + 1):
    v = ws_data.cell(row=ri, column=1).value
    if v is None or v == "":
        rows_to_delete.append(ri)

print(f"Rows with empty Raman Shift (A): {len(rows_to_delete)} -> deleting bottom-up")
for ri in reversed(rows_to_delete):
    ws_data.delete_rows(ri)

# ---- Step 4: Delete header row (data starts from row 1) ----
ws_data.delete_rows(1)
print(f"Header row deleted. Final data: {ws_data.max_row} rows x {ws_data.max_column} columns")

# ---- Step 5: Save ----
saved = False
alt_path = None
for attempt in range(2):
    try:
        wb.save(XLSX_PATH)
        print(f"\n[OK] Saved: {XLSX_PATH}")
        saved = True
        break
    except PermissionError:
        if attempt == 0:
            alt = FOLDER_NAME + "_已修改.xlsx"  # _已修改
            alt_path = os.path.join(BASE_DIR, alt)
            print(f"[WARN] File locked, saving as: {alt}")
            wb.save(alt_path)
            saved = True
            break
    except Exception as exc:
        if attempt == 0:
            alt = FOLDER_NAME + "_已修改.xlsx"
            alt_path = os.path.join(BASE_DIR, alt)
            print(f"[WARN] Save error ({exc}), fallback: {alt}")
            wb.save(alt_path)
            saved = True
            break

if not saved:
    print("[ERR] Save failed. Check if the file is open in Excel.")
    sys.exit(1)

print("\n=== Phase 1 Complete ===")
